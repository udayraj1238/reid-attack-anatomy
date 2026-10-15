"""Generate a 16:9 widescreen PowerPoint presentation deck from real experimental results.
Outputs: docs/presentation_deck.pptx
"""
import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULTS = os.path.join(ROOT, "results")
DOCS = os.path.join(ROOT, "docs")
OUT_PATH = os.path.join(DOCS, "presentation_deck.pptx")

# Color palette (Modern dark academic slate)
COLOR_BG = RGBColor(15, 23, 42)        # #0F172A Dark Navy
COLOR_CARD = RGBColor(30, 41, 59)      # #1E293B Slate Card
COLOR_BORDER = RGBColor(51, 65, 85)    # #334155
COLOR_ACCENT = RGBColor(56, 189, 248)  # #38BDF8 Cyan
COLOR_ACCENT_ALT = RGBColor(129, 140, 248) # #818CF8 Indigo
COLOR_WHITE = RGBColor(248, 250, 252) # #F8FAFC
COLOR_MUTED = RGBColor(148, 163, 184) # #94A3B8
COLOR_GREEN = RGBColor(52, 211, 153)  # #34D399 Emerald
COLOR_RED = RGBColor(248, 113, 113)   # #F87171 Red

def set_slide_bg(slide):
    # Add full slide dark rectangle
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category="RE-ID ADVERSARIAL ANATOMY"):
    # Category tag
    tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    p_cat = tb_cat.text_frame.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_ACCENT
    
    # Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Load real facts if available
    facts_file = os.path.join(RESULTS, "facts.json")
    facts = json.load(open(facts_file)) if os.path.exists(facts_file) else {}

    # -------------------------------------------------------------------------
    # Slide 1: Title
    # -------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)
    
    add_card(s1, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1), COLOR_CARD, COLOR_ACCENT)
    
    tb = s1.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.333), Inches(2.2))
    p = tb.text_frame.paragraphs[0]
    p.text = "Anatomy of an Attack on Person Re-Identification"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    
    p2 = tb.text_frame.add_paragraph()
    p2.text = "(+ An EOT Universal Invisibility Torso Patch & Adversarial Defense)"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT
    
    tb_meta = s1.shapes.add_textbox(Inches(2.0), Inches(4.3), Inches(9.333), Inches(1.5))
    pm = tb_meta.text_frame.paragraphs[0]
    pm.text = "Deep Learning Project • Empirical Benchmark on Market-1501"
    pm.font.size = Pt(14)
    pm.font.color.rgb = COLOR_MUTED
    
    pm2 = tb_meta.text_frame.add_paragraph()
    pm2.text = "Architectures: ResNet-50 (BoT), DenseNet-121, ViT-Small/16 | 100% Empirically Measured"
    pm2.font.size = Pt(13)
    pm2.font.color.rgb = COLOR_WHITE

    # -------------------------------------------------------------------------
    # Slide 2: Why This Matters (The Problem Formulation)
    # -------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "Why Re-ID Security Differs from Standard Classification")
    
    c1 = add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb1 = s2.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.3))
    p = tb1.text_frame.paragraphs[0]
    p.text = "Open-Set Ranking vs Fixed Labels"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT
    bullets = [
        "Closed-Set Classification: Predicts 1 of N known training classes using softmax logits.",
        "Person Re-ID: Queries and gallery belong to unseen identities across non-overlapping surveillance cameras.",
        "Scoring Metric: Evaluated on gallery distance ranking (Rank-1, Rank-5, mAP), not class accuracy.",
        "Silent Retrieval Failure: An attacker doesn't need to force a label—simply degrading Rank-1 displaces true matches out of the top-k results."
    ]
    for b in bullets:
        p_b = tb1.text_frame.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_WHITE
        p_b.space_before = Pt(10)
        
    c2 = add_card(s2, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    tb2 = s2.shapes.add_textbox(Inches(7.2), Inches(2.0), Inches(5.0), Inches(4.3))
    p = tb2.text_frame.paragraphs[0]
    p.text = "Core Research Questions We Answer"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT_ALT
    questions = [
        "1. Embedding vs Logit Attacks: Does directly pushing embeddings in metric space outperform classical classification-head attacks?",
        "2. Cross-Architecture Transfer: Do adversarial directions crafted on CNNs transfer to other CNNs and Vision Transformers?",
        "3. Universal Physical Feasibility: Can a single fixed torso patch fool multiple architectures without per-image optimization?",
        "4. Adversarial Training Cost: What is the exact clean accuracy trade-off when hardening Re-ID with PGD adversarial training?"
    ]
    for q in questions:
        p_q = tb2.text_frame.add_paragraph()
        p_q.text = q
        p_q.font.size = Pt(13)
        p_q.font.color.rgb = COLOR_WHITE
        p_q.space_before = Pt(10)

    # -------------------------------------------------------------------------
    # Slide 3: Clean Baselines & Architecture Details
    # -------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Victim Models: Architecture & Clean Performance (Market-1501)")
    
    # Table of baselines
    table_shape = s3.shapes.add_table(4, 6, Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.4))
    table = table_shape.table
    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(1.8)
    table.columns[3].width = Inches(1.8)
    table.columns[4].width = Inches(1.8)
    table.columns[5].width = Inches(1.9)
    
    headers = ["Architecture", "Parameters", "Rank-1 (%)", "Rank-5 (%)", "Rank-10 (%)", "mAP (%)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid(); cell.fill.fore_color.rgb = COLOR_BORDER
        p = cell.text_frame.paragraphs[0]
        p.text = h; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = COLOR_ACCENT
        p.alignment = PP_ALIGN.CENTER
        
    data = [
        ["ResNet-50 (BoT / BNNeck)", "25.1 M", "86.97 %", "95.40 %", "97.36 %", "70.64 %"],
        ["DenseNet-121", "7.7 M", "77.35 %", "91.95 %", "95.07 %", "56.63 %"],
        ["ViT-Small / 16 (DeiT)", "21.9 M", "91.75 %", "96.67 %", "98.01 %", "79.69 %"]
    ]
    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid(); cell.fill.fore_color.rgb = COLOR_CARD
            p = cell.text_frame.paragraphs[0]
            p.text = val; p.font.size = Pt(13); p.font.color.rgb = COLOR_WHITE
            p.alignment = PP_ALIGN.CENTER
            
    # Technical training note card
    c_note = add_card(s3, Inches(0.8), Inches(4.6), Inches(11.7), Inches(2.1))
    tb_note = s3.shapes.add_textbox(Inches(1.1), Inches(4.7), Inches(11.1), Inches(1.9))
    pn = tb_note.text_frame.paragraphs[0]
    pn.text = "Standard Training Recipe & Protocol Constraints"
    pn.font.size = Pt(15); pn.font.bold = True; pn.font.color.rgb = COLOR_ACCENT
    bullets_recipe = [
        "Initialization: All backbones use ImageNet-1K pretrained weights fine-tuned on Market-1501 (12,936 training images, 751 identities).",
        "Loss Formulation: Joint ID cross-entropy loss + Batch-Hard Triplet loss with margin m = 0.3 (PK sampler: P=16 identities, K=4 instances).",
        "BNNeck Head: Batch normalization layer separates metric embedding f (normalized for cosine ranking) from classification logits.",
        "Market-1501 Protocol: Evaluated on 3,368 queries vs 15,913 gallery targets across 6 cameras. Junk rule strictly enforced (same PID + camera ignored)."
    ]
    for br in bullets_recipe:
        p_br = tb_note.text_frame.add_paragraph()
        p_br.text = "• " + br
        p_br.font.size = Pt(12)
        p_br.font.color.rgb = COLOR_WHITE
        p_br.space_before = Pt(4)

    # -------------------------------------------------------------------------
    # Slide 4: Attack Formulations & Mathematical Objectives
    # -------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "Threat Model & Objective Formulations in Metric Space")
    
    # 3 Objective cards
    obj_defs = [
        ("Evasion Objective", "Pushes the adversarial embedding as far away from its clean representation as possible.",
         "L_evade = - ( f(x_adv) · f(x) )",
         "Minimizes cosine similarity between perturbed and clean feature vectors. Forces true matches out of the top gallery ranks.",
         COLOR_ACCENT),
        ("Impersonation Objective", "Pulls the query embedding towards the centroid embedding of an unauthorized identity.",
         "L_impersonate = + ( f(x_adv) · e_target )",
         "Maximizes cosine similarity to target centroid. Measured by Hit@1: target identity returned as Rank-1 match.",
         COLOR_ACCENT_ALT),
        ("Classification Logit Baseline", "Standard adversarial classification attack on the training ID-classifier head.",
         "L_logit = CrossEntropy( W · f(x_adv), y_true )",
         "Maximizes cross-entropy loss on training classes. Used to test whether classification attacks suffice for open-set Re-ID.",
         COLOR_MUTED)
    ]
    for i, (title, desc, formula, impact, col) in enumerate(obj_defs):
        left = Inches(0.8 + i * 4.0)
        c = add_card(s4, left, Inches(1.8), Inches(3.7), Inches(5.0))
        tb = s4.shapes.add_textbox(left + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.5))
        p = tb.text_frame.paragraphs[0]
        p.text = title; p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = col
        
        pf = tb.text_frame.add_paragraph()
        pf.text = formula; pf.font.size = Pt(12); pf.font.bold = True; pf.font.color.rgb = COLOR_WHITE
        pf.space_before = Pt(12)
        
        pd = tb.text_frame.add_paragraph()
        pd.text = desc; pd.font.size = Pt(12); pd.font.color.rgb = COLOR_MUTED
        pd.space_before = Pt(10)
        
        pi = tb.text_frame.add_paragraph()
        pi.text = impact; pi.font.size = Pt(12); pi.font.color.rgb = COLOR_WHITE
        pi.space_before = Pt(10)

    # -------------------------------------------------------------------------
    # Slide 5: White-Box Evasion on ResNet-50
    # -------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "White-Box Evasion: Rank-1 Collapses to 0.0% at eps = 8/255")
    
    # Left: Stats Card
    c_left = add_card(s5, Inches(0.8), Inches(1.8), Inches(5.0), Inches(5.1))
    tb_l = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.6), Inches(4.6))
    pl = tb_l.text_frame.paragraphs[0]
    pl.text = "Epsilon Degradation (ResNet-50)"
    pl.font.size = Pt(17); pl.font.bold = True; pl.font.color.rgb = COLOR_ACCENT
    
    stats_text = [
        "Clean Baseline: 86.97 % Rank-1 | 70.64 % mAP",
        "FGSM (1-step eps=8): Drops to 39.58 % Rank-1",
        "PGD-10 (eps=8): Complete collapse to 0.00 % Rank-1 (0.05 % mAP)",
        "MI-FGSM (eps=8): Complete collapse to 0.00 % Rank-1 (0.05 % mAP)",
        "Logit-PGD (eps=8): 0.30 % Rank-1 (Embedding PGD is strictly stronger)",
        "Even at eps = 2/255 (barely perceptible): PGD drives Rank-1 to 0.12 %!"
    ]
    for st in stats_text:
        p_st = tb_l.text_frame.add_paragraph()
        p_st.text = "• " + st
        p_st.font.size = Pt(12)
        p_st.font.color.rgb = COLOR_WHITE
        p_st.space_before = Pt(8)
        
    # Right: Add figure if exists
    img_eps = os.path.join(RESULTS, "eps_degradation_curve.png")
    if os.path.exists(img_eps):
        add_card(s5, Inches(6.1), Inches(1.8), Inches(6.4), Inches(5.1))
        s5.shapes.add_picture(img_eps, Inches(6.3), Inches(2.0), Inches(6.0), Inches(4.6))

    # -------------------------------------------------------------------------
    # Slide 6: Visual Evidence (Retrieval Before/After & Grad-CAM)
    # -------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "Visualizing Failure: Retrieval Before/After & Attention Drift")
    
    img_ret = os.path.join(RESULTS, "retrieval_before_after.png")
    img_cam = os.path.join(RESULTS, "gradcam_analysis.png")
    
    if os.path.exists(img_ret):
        add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.4))
        s6.shapes.add_picture(img_ret, Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.0))
        tb = s6.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(5.6), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text = "Query Retrieval Ranking: Clean top-5 matches true ID; under PGD, true match drops entirely from top results."
        p.font.size = Pt(11); p.font.color.rgb = COLOR_MUTED
        
    if os.path.exists(img_cam):
        add_card(s6, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.4))
        s6.shapes.add_picture(img_cam, Inches(7.1), Inches(2.0), Inches(5.2), Inches(4.0))
        tb = s6.shapes.add_textbox(Inches(6.9), Inches(6.3), Inches(5.6), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text = "Grad-CAM Feature Saliency: Gradient perturbations shift network focus away from person torso to background clutter."
        p.font.size = Pt(11); p.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------------------
    # Slide 7: Cross-Architecture Transfer Matrix
    # -------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Cross-Architecture Transfer Matrix: Asymmetry Between CNNs & ViT")
    
    img_mat = os.path.join(RESULTS, "transfer_matrix_pgd_eps8_rank1.png")
    if os.path.exists(img_mat):
        add_card(s7, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.1))
        s7.shapes.add_picture(img_mat, Inches(1.0), Inches(2.0), Inches(5.6), Inches(4.6))
        
    c_right = add_card(s7, Inches(7.1), Inches(1.8), Inches(5.4), Inches(5.1))
    tb_r = s7.shapes.add_textbox(Inches(7.3), Inches(2.0), Inches(5.0), Inches(4.6))
    pr = tb_r.text_frame.paragraphs[0]
    pr.text = "Key Empirical Transfer Observations"
    pr.font.size = Pt(17); pr.font.bold = True; pr.font.color.rgb = COLOR_ACCENT
    
    insights = [
        "White-Box vs Black-Box: Mean Rank-1 is 0.0 % white-box vs 61.04 % under black-box PGD transfer.",
        "CNN-to-CNN Transfer: Attacks transfer moderately between CNN architectures: ResNet-50 PGD drops DenseNet-121 from 77.35 % to 30.46 % (MI-FGSM to 17.81 %).",
        "CNN-to-ViT Resistance: Perturbations crafted on ResNet-50 or DenseNet-121 barely hurt ViT-Small (Rank-1 remains 88.15 % and 88.66 % vs 91.75 % clean).",
        "Hypothesis: Local convolutional receptive fields share inductive bias, whereas ViT's global self-attention produces divergent gradient landscapes.",
        "MI-FGSM Advantage: Momentum smooths gradient oscillations and reduces mean black-box Rank-1 from 61.04 % (PGD) to 50.64 %."
    ]
    for ins in insights:
        p_ins = tb_r.text_frame.add_paragraph()
        p_ins.text = "• " + ins
        p_ins.font.size = Pt(12)
        p_ins.font.color.rgb = COLOR_WHITE
        p_ins.space_before = Pt(8)

    # -------------------------------------------------------------------------
    # Slide 8: Universal Invisibility Torso Patch (EOT)
    # -------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Universal Torso Patch: One Printable Pattern Evades Retrieval")
    
    img_patch_eff = os.path.join(RESULTS, "patch_effect.png")
    img_patch_ex = os.path.join(RESULTS, "patch_examples.png")
    
    if os.path.exists(img_patch_ex):
        add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.0))
        s8.shapes.add_picture(img_patch_ex, Inches(1.0), Inches(1.9), Inches(5.2), Inches(3.7))
        
    c_pstats = add_card(s8, Inches(6.7), Inches(1.8), Inches(5.8), Inches(5.1))
    tbp = s8.shapes.add_textbox(Inches(6.9), Inches(2.0), Inches(5.4), Inches(4.7))
    pt = tbp.text_frame.paragraphs[0]
    pt.text = "Torso Patch Performance on ResNet-50"
    pt.font.size = Pt(17); pt.font.bold = True; pt.font.color.rgb = COLOR_ACCENT
    
    patch_pts = [
        "Control Experiment (Random Noise Patch): Leaves Rank-1 at 73.19 % (Rank-1 drop is only 13.78 % due to occlusion).",
        "Trained Universal Patch (64x48): Collapses Rank-1 to 2.67 % (mAP 3.29 %) — an 84.30 % absolute drop!",
        "EOT Robustness: Under random testing scale, rotation, translation, and noise, Rank-1 remains crushed at 2.73 %.",
        "Cross-Model Transfer: The patch trained only on ResNet-50 drops DenseNet-121 Rank-1 to 13.63 % without retraining.",
        "EOT Training Objective: Minimizes cosine similarity to clean features over 40 epochs with TV smoothness regularization."
    ]
    for pt_item in patch_pts:
        p_pi = tbp.text_frame.add_paragraph()
        p_pi.text = "• " + pt_item
        p_pi.font.size = Pt(12)
        p_pi.font.color.rgb = COLOR_WHITE
        p_pi.space_before = Pt(8)
        
    if os.path.exists(img_patch_eff):
        add_card(s8, Inches(0.8), Inches(6.0), Inches(5.6), Inches(1.0))
        tbc = s8.shapes.add_textbox(Inches(0.9), Inches(6.05), Inches(5.4), Inches(0.9))
        p = tbc.text_frame.paragraphs[0]
        p.text = "Trained on 40 epochs with Total Variation (TV) loss to prevent high-frequency noise and ensure physical printability."
        p.font.size = Pt(11); p.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------------------
    # Slide 9: Adversarial Training Defense (PGD eps=8)
    # -------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Defense Benchmark: Adversarial Training Trade-Off")
    
    img_def = os.path.join(RESULTS, "defense_tradeoff.png")
    if os.path.exists(img_def):
        add_card(s9, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.1))
        s9.shapes.add_picture(img_def, Inches(1.0), Inches(2.0), Inches(5.6), Inches(4.6))
        
    c_def = add_card(s9, Inches(7.1), Inches(1.8), Inches(5.4), Inches(5.1))
    tb_d = s9.shapes.add_textbox(Inches(7.3), Inches(2.0), Inches(5.0), Inches(4.6))
    pd = tb_d.text_frame.paragraphs[0]
    pd.text = "Empirical Defense Evaluation (ResNet-50)"
    pd.font.size = Pt(17); pd.font.bold = True; pd.font.color.rgb = COLOR_ACCENT
    
    def_bullets = [
        "White-Box Robustness Gain: Robust Rank-1 rises from 0.00 % (standard model) to 29.19 % (robust model) under white-box PGD-10.",
        "Clean Accuracy Trade-off: Clean Rank-1 drops from 86.97 % to 74.52 % (a 12.45 % penalty), consistent with classical robust classification trade-offs.",
        "Black-Box Protection: When attacked by perturbations from the standard ResNet-50, the adversarially trained model retains 72.39 % Rank-1 (mAP 49.96 %).",
        "Caveat on Evaluation: The defense was trained against PGD-5 and evaluated against PGD-10; adaptive attacks (e.g. PGD-50) would likely erode robustness further."
    ]
    for db in def_bullets:
        p_db = tb_d.text_frame.add_paragraph()
        p_db.text = "• " + db
        p_db.font.size = Pt(12)
        p_db.font.color.rgb = COLOR_WHITE
        p_db.space_before = Pt(8)

    # -------------------------------------------------------------------------
    # Slide 10: Conclusion & Honest Scientific Limitations
    # -------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "Summary, Limitations & Future Directions")
    
    c_sum = add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    tb_s = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    ps = tb_s.text_frame.paragraphs[0]
    ps.text = "Core Takeaways"
    ps.font.size = Pt(17); ps.font.bold = True; ps.font.color.rgb = COLOR_GREEN
    concl = [
        "Re-ID is uniquely vulnerable: Gradient evasion attacks drive Rank-1 retrieval to 0.00 % without requiring classification label access.",
        "Architecture Matters: ViTs exhibit substantial intrinsic resistance to CNN-crafted transfer attacks (maintaining >88 % Rank-1).",
        "Physical Feasibility: A single 64x48 torso patch trained with EOT effectively blinds surveillance models (Rank-1 86.97 % -> 2.67 %).",
        "Defense Reality: Adversarial training restores robust Rank-1 to 29.19 % but incurs a 12.45 % clean retrieval penalty."
    ]
    for cc in concl:
        p_cc = tb_s.text_frame.add_paragraph()
        p_cc.text = "✔ " + cc
        p_cc.font.size = Pt(12)
        p_cc.font.color.rgb = COLOR_WHITE
        p_cc.space_before = Pt(8)
        
    c_lim = add_card(s10, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1))
    tb_lim = s10.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.2), Inches(4.6))
    plim = tb_lim.text_frame.paragraphs[0]
    plim.text = "Scientific Transparency & Limitations"
    plim.font.size = Pt(17); plim.font.bold = True; plim.font.color.rgb = COLOR_RED
    lims = [
        "Backbones Pretrained: Backbones use ImageNet-1K pretrained weights fine-tuned on Market-1501, not trained from scratch.",
        "Epoch Definition: One 'epoch' represents one pass over 751 identities (4 images each), approximately 23 % of total dataset images.",
        "Digital-Only Patch: Torso patch was evaluated via digital transformation overlay; physical printing on clothing was not experimentally validated.",
        "Single Seed Evaluation: Experiments conducted with fixed seed (seed=0); multi-seed confidence intervals remain future work.",
        "Adaptive Attacks: Defense evaluated on PGD-10; advanced AutoAttack/PGD-50 evaluations are needed for definitive robustness bounds."
    ]
    for lm in lims:
        p_lm = tb_lim.text_frame.add_paragraph()
        p_lm.text = "• " + lm
        p_lm.font.size = Pt(12)
        p_lm.font.color.rgb = COLOR_WHITE
        p_lm.space_before = Pt(8)

    prs.save(OUT_PATH)
    print(f"Presentation saved successfully to: {OUT_PATH}")

if __name__ == "__main__":
    main()
